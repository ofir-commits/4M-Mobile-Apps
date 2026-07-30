# -*- coding: utf-8 -*-
"""Write cached results into the summary sheet's formula cells.

recalc.py cannot run here: LibreOffice in this sandbox fails to load even a
two-cell openpyxl workbook ("source file could not be loaded"), so it can
neither compute nor cache anything. Without cached values every COUNTIF cell
reads back as None to pandas, to load_workbook(data_only=True) and to preview
panes — the numbers would look empty until someone opened the file in Excel.

So the formulas stay (a מורשה editing an F cell must see the summary follow),
and the value each one evaluates to is computed here in Python and injected into
the sheet XML as the cached <v>. fullCalcOnLoad is also set, so Excel and
LibreOffice recompute from the formulas on open and would overwrite anything
wrong here rather than trust it.
"""
import re
import shutil
import os
import sys
import zipfile

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from assess import TEMPLATES, OK, NO, NA

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'anshilo-accessibility-form.xlsx')
FIRST, LAST = 4, 45
FIRST_ROW, TOTAL_ROW = 16, 27
SUMMARY = 'סיכום ומקרא'

# ---- compute what each formula evaluates to, from the sheets themselves ----
wb = openpyxl.load_workbook(SRC)
values = {}                       # 'B16' -> number
cols = {'B': OK, 'C': NO, 'D': NA}

for n, (sheet_name, _, _) in enumerate(TEMPLATES):
    ws = wb[sheet_name]
    got = [ws.cell(row=r, column=6).value for r in range(FIRST, LAST + 1)]
    row = FIRST_ROW + n
    for col, want in cols.items():
        values['%s%d' % (col, row)] = sum(1 for v in got if v == want)
    values['E%d' % row] = sum(1 for v in got if v is None or str(v).strip() == '')
    values['F%d' % row] = len(got)

for col in 'BCDEF':
    values['%s%d' % (col, TOTAL_ROW)] = sum(
        values['%s%d' % (col, FIRST_ROW + n)] for n in range(len(TEMPLATES)))

# ---- set fullCalcOnLoad so a real spreadsheet app recomputes anyway ----
wb.calculation.fullCalcOnLoad = True
wb.save(SRC)

# ---- find which XML part is the summary sheet ----
wb2 = openpyxl.load_workbook(SRC)
sheet_index = wb2.sheetnames.index(SUMMARY)
target = None
with zipfile.ZipFile(SRC) as z:
    wbxml = z.read('xl/workbook.xml').decode('utf-8')
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    rid = re.findall(r'<sheet[^>]*r:id="([^"]+)"', wbxml)[sheet_index]
    # Attribute order in the rels part is not fixed — locate the element by Id,
    # then read Target out of it, rather than assuming Id comes first.
    element = next(e for e in re.findall(r'<Relationship\b[^>]*/?>', rels)
                   if 'Id="%s"' % rid in e)
    path = re.search(r'Target="([^"]+)"', element).group(1)
    target = path.lstrip('/') if path.startswith('/xl/') else 'xl/' + path.lstrip('/')
print('summary sheet part:', target)

# ---- inject the cached value after each </f> ----
injected = []


def add_value(m):
    ref, attrs, formula = m.group('ref'), m.group('attrs'), m.group('f')
    if ref not in values:
        return m.group(0)
    injected.append(ref)
    return '<c r="%s"%s><f>%s</f><v>%d</v></c>' % (ref, attrs, formula, values[ref])


# openpyxl already emits an empty <v /> placeholder after the formula, so the
# pattern has to tolerate that as well as its absence.
CELL = re.compile(
    r'<c r="(?P<ref>[A-Z]+\d+)"(?P<attrs>[^>]*)>\s*<f>(?P<f>[^<]*)</f>'
    r'(?:\s*<v\s*/>|\s*<v>[^<]*</v>)?\s*</c>')

with zipfile.ZipFile(SRC) as zin:
    names = zin.namelist()
    parts = {n: zin.read(n) for n in names}

xml = parts[target].decode('utf-8')
xml = CELL.sub(add_value, xml)
parts[target] = xml.encode('utf-8')

OUT = SRC + '.tmp'
with zipfile.ZipFile(OUT, 'w', zipfile.ZIP_DEFLATED) as zout:
    for n in names:
        zout.writestr(n, parts[n])
shutil.move(OUT, SRC)

print('cached %d formula cells' % len(injected))
missing = sorted(set(values) - set(injected))
if missing:
    print('WARNING: no formula found at', missing)
    sys.exit(1)
